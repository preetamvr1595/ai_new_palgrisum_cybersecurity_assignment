import docx
from typing import Dict, Any, Tuple

def extract_docx(file_path: str) -> Tuple[str, Dict[str, Any]]:
    """
    Extracts text and metadata from a DOCX file using python-docx.
    """
    doc = docx.Document(file_path)
    text_content = []
    
    metadata = {
        "author": doc.core_properties.author,
        "title": doc.core_properties.title,
        "created": doc.core_properties.created,
        "modified": doc.core_properties.modified,
        "revision": doc.core_properties.revision
    }

    for para in doc.paragraphs:
        if para.text.strip():
            text_content.append(para.text)

    # Extract text from tables as well
    for table in doc.tables:
        for row in table.rows:
            row_data = []
            for cell in row.cells:
                if cell.text.strip():
                    row_data.append(cell.text.strip())
            if row_data:
                text_content.append(" | ".join(row_data))

    return "\n".join(text_content), metadata
